import io, copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SCR = sys.argv[1]
OUT = sys.argv[2]
tpl = Presentation(f"{SCR}/tpl.pptx")
old = Presentation(f"{SCR}/tula.pptx")

NAVY = "1F3F6E"; BLUE = "0070C0"; GREEN = "1E7B4A"; ORANGE = "D9661F"; PURPLE = "6B3FA0"
INK = "1F2937"; GREY = "5B6573"; LIGHT = "F1F4F8"; LBLUE = "E6EEF8"; LGREEN = "E6F4EC"
LPURPLE = "F0EAF8"; LORANGE = "FCEFE5"; WHITE = "FFFFFF"; LINE = "C9D3E0"
FONT = "Calibri"

def rgb(h): return RGBColor.from_string(h)

# ---------- helpers ----------
def _fmt_run(r, size, bold=False, color=INK, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color); r.font.name = FONT

def text(sl, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.0, italic=False, spacing=0):
    """paras: str | list of (str | list of (text, {opts}))"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    m = Inches(margin); tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    if isinstance(paras, str): paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing: para.space_after = Pt(spacing)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run(); r.text = t
            _fmt_run(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", italic))
    return tb

def box(sl, x, y, w, h, fill=LIGHT, line=None, paras=None, size=11, color=INK, bold=False,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12, lw=1.0, dash=False, margin=0.06):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.shadow.inherit = False
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            s.line.dash_style = MSO_LINE.DASH
    else: s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    m = Inches(margin); tf.margin_left = tf.margin_right = m; tf.margin_top = tf.margin_bottom = Inches(0.03)
    if paras is not None:
        if isinstance(paras, str): paras = [paras]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            runs = p if isinstance(p, list) else [(p, {})]
            for t, o in runs:
                r = para.add_run(); r.text = t
                _fmt_run(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", False))
    return s

def arrow(sl, x1, y1, x2, y2, color=BLUE, w=1.5, head=True, dash=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("h", "med")
    return c

def label(sl, x, y, w, t, color=NAVY, size=12.5):
    return text(sl, x, y, w, 0.3, [[(t, {"bold": True, "color": color, "size": size})]])

def circle(sl, x, y, d, t, fill=NAVY, size=10):
    return box(sl, x, y, d, d, fill=fill, paras=[[(t, {"bold": True, "color": WHITE, "size": size})]],
               shape=MSO_SHAPE.OVAL, margin=0)

def pic_blob(slide_idx, name_contains):
    for sh in old.slides[slide_idx].shapes:
        if sh.shape_type == 13 and any(name_contains.lower() in d.lower() for d in sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr")):
            return sh.image.blob
    for sh in old.slides[slide_idx].shapes:
        if sh.shape_type == 13 and name_contains in (sh.image.filename or "") + sh.name:
            return sh.image.blob
    for sh in old.slides[slide_idx].shapes:
        if sh.shape_type == 13 and name_contains.lower() in sh.name.lower():
            return sh.image.blob
    return None

def set_title(sl, t, size=None):
    for sh in sl.shapes:
        if sh.is_placeholder and sh.placeholder_format.type is not None and "Title" in sh.name:
            p = sh.text_frame.paragraphs[0]
            runs = p.runs
            runs[0].text = t
            if size: runs[0].font.size = Pt(size)
            for r in runs[1:]: r.text = ""
            for extra in sh.text_frame.paragraphs[1:]:
                for r in extra.runs: r.text = ""
            return sh

def team_oval(sl):
    for sh in sl.shapes:
        if sh.has_text_frame and "Team Name" in sh.text_frame.text:
            tf = sh.text_frame
            p0 = tf.paragraphs[0]
            proto = p0.runs[0]
            proto.text = "Team"; _fmt_run(proto, 10, False, NAVY)
            for r in p0.runs[1:]: r.text = ""
            p1 = tf.add_paragraph(); p1.alignment = p0.alignment
            r = p1.add_run(); r.text = "FriendlyFire"
            _fmt_run(r, 10, True, NAVY)
            for extra in tf.paragraphs[1:-1]:
                for rr in extra.runs: rr.text = ""

def footer(sl):
    for sh in sl.shapes:
        if sh.has_text_frame and "@SIH Idea submission" in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("Template", "TULA")

def remove(sl, name):
    for sh in list(sl.shapes):
        if sh.name == name:
            sh._element.getparent().remove(sh._element)

# ---------- drop instructions slide (7) ----------
sldIdLst = tpl.slides._sldIdLst
last = sldIdLst[6]
tpl.part.drop_rel(last.rId); sldIdLst.remove(last)

S = tpl.slides

# ================= SLIDE 1: TITLE =================
s = S[0]
remove(s, "TextBox 9")
box(s, 0.42, 2.0, 6.35, 0.8, fill=NAVY, paras=[[("TULA  ", {"bold": True, "size": 26, "color": WHITE}),
    ("Unified Digital Lifecycle Platform for Legal Metrology", {"bold": True, "size": 13, "color": WHITE})]],
    align=PP_ALIGN.LEFT, margin=0.2, radius=0.08)
items = [("Problem Statement ID", "26036"),
         ("Problem Statement Title", "Development of an Online Verification System for Weighing and Measuring Instruments"),
         ("Theme", "Miscellaneous"), ("PS Category", "Software"),
         ("Team ID", "_______"), ("Team Name", "FriendlyFire")]
tb = s.shapes.add_textbox(Inches(0.42), Inches(3.0), Inches(6.4), Inches(3.3))
tf = tb.text_frame; tf.word_wrap = True
for i, (k, v) in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr(); pPr.set("marL", str(Inches(0.3))); pPr.set("indent", str(-Inches(0.3)))
    bu = etree.SubElement(pPr, qn("a:buFont")); bu.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "•")
    r = p.add_run(); r.text = f"{k} – "; _fmt_run(r, 17, True, INK)
    r = p.add_run(); r.text = v; _fmt_run(r, 17, False, INK)
text(s, 0.42, 6.45, 6.6, 0.4, [[("One instrument → one digital ID → one certificate anyone can check",
     {"italic": True, "bold": True, "color": BLUE, "size": 12.5})]])
s.notes_slide.notes_text_frame.text = (
    "Open with the problem in one line: verification of scales, fuel pumps and weighbridges is still paper and siloed, "
    "and a shopper has no way to tell a genuine stamp from a fake. TULA gives every instrument one permanent digital ID "
    "and runs its whole verification life on it. Fill in the Team ID before exporting to PDF.")

# ================= SLIDE 2: IDEA =================
s = S[1]
set_title(s, "TULA: ONE INSTRUMENT, ONE DIGITAL ID", size=28)
team_oval(s); footer(s); remove(s, "TextBox 8")
L, R = 0.4, 12.93; W = R - L
label(s, L, 1.28, 9, "PROPOSED SOLUTION")
box(s, L, 1.58, W, 0.52, fill=LBLUE, align=PP_ALIGN.LEFT, margin=0.15, radius=0.06, paras=[[
    ("A web + mobile platform where every weighing and measuring instrument gets a permanent Digital ID. ", {"bold": True, "size": 12.5, "color": NAVY}),
    ("Application, scrutiny, inspection, certificate, re-verification and enforcement all attach to that one ID.", {"size": 12.5, "color": INK})]])

label(s, L, 2.2, 9, "DETAILED EXPLANATION: the life of one instrument in TULA")
steps = [("Register", "Digital ID issued", BLUE), ("Apply", "online, fee auto-calc", BLUE),
         ("Scrutiny", "rules + AI copilot", PURPLE), ("Assign", "district LMO / GATC", BLUE),
         ("Inspect", "offline field app", BLUE), ("Decide", "readings vs MPE", ORANGE),
         ("Certificate", "signed QR + PDF", GREEN), ("Verify", "anyone scans, no login", GREEN)]
bw, gap, by, bh = 1.38, (W - 8 * 1.38) / 7, 2.62, 0.8
for i, (t, sub, c) in enumerate(steps):
    x = L + i * (bw + gap)
    fill = {BLUE: WHITE, PURPLE: LPURPLE, ORANGE: LORANGE, GREEN: LGREEN}[c]
    box(s, x, by, bw, bh, fill=fill, line=c, lw=1.5, radius=0.08,
        paras=[[(t, {"bold": True, "size": 13, "color": c if c != BLUE else NAVY})], [(sub, {"size": 9.5, "color": GREY})]])
    circle(s, x - 0.1, by - 0.13, 0.3, str(i + 1), fill=c if c != BLUE else NAVY, size=9.5)
    if i < 7:
        arrow(s, x + bw + 0.02, by + bh / 2, x + bw + gap - 0.02, by + bh / 2, color=NAVY, w=1.5)
# expiry loop below the flow: Certificate(7) -> Apply(2)
x7 = L + 6 * (bw + gap) + bw / 2; x2 = L + 1 * (bw + gap) + bw / 2; ly = by + bh + 0.22
arrow(s, x7, by + bh, x7, ly, color=GREEN, w=1.75, head=False)
arrow(s, x7, ly, x2, ly, color=GREEN, w=1.75, head=False)
arrow(s, x2, ly, x2, by + bh + 0.02, color=GREEN, w=1.75)
text(s, x2 + 0.6, ly + 0.02, 6.5, 0.28, [[("Expiry alerts at 30 / 15 / 7 days  →  re-verification on the same Digital ID",
     {"bold": True, "size": 11, "color": GREEN})]])
# fail branch
xd = L + 5 * (bw + gap)
text(s, xd + bw + 0.05, by + bh + 0.02, 0, 0, "")

label(s, L, 4.02, 9, "HOW IT ADDRESSES THE PROBLEM")
probs = [("Records on paper or isolated local systems", "One central record per instrument"),
         ("Manual scheduling and slow pendency", "Rule-based fee, routing and live queues"),
         ("Consumers can't check a stamp", "Scan the QR, see validity instantly"),
         ("Re-verification dates get missed", "Automatic reminders before expiry")]
cw = (W - 3 * 0.2) / 4
for i, (p, a) in enumerate(probs):
    x = L + i * (cw + 0.2)
    box(s, x, 4.32, cw, 0.9, fill=LIGHT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.14, radius=0.06,
        paras=[[(p, {"size": 10.5, "color": GREY})], [("→ " + a, {"bold": True, "size": 12, "color": NAVY})]])

label(s, L, 5.33, 9, "INNOVATION AND UNIQUENESS")
inn = [("Instrument-centric", "One ID carries the full history across owners and years", NAVY),
       ("Signed QR", "Ed25519 signature: a copied or edited certificate fails the scan", GREEN),
       ("Rules as data", "Each state sets fees, validity and checklists without code changes", BLUE),
       ("Gemini copilot", "Reads documents and nameplates, flags mismatches. Officer decides", PURPLE),
       ("Citizen as auditor", "Any shopper can scan and report a bad scale in two taps", ORANGE)]
tw = (W - 4 * 0.15) / 5
for i, (t, d, c) in enumerate(inn):
    x = L + i * (tw + 0.15)
    box(s, x, 5.63, tw, 1.17, fill=c, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.12, radius=0.07,
        paras=[[(t, {"bold": True, "size": 12.5, "color": WHITE})], [(d, {"size": 10, "color": WHITE})]])
s.notes_slide.notes_text_frame.text = (
    "Walk the 8 steps left to right using one instrument, e.g. a 100-tonne weighbridge in Ahmedabad. "
    "Stress: every event lands on the same Digital ID, so history is never lost. Gemini only assists scrutiny; "
    "fees, MPE pass/fail and the final decision are deterministic or human. The green loop is the re-verification cycle.")

# ================= SLIDE 3: TECHNICAL =================
s = S[2]
team_oval(s); footer(s); remove(s, "TextBox 8")
LX, LW = 0.4, 8.3
label(s, LX, 1.28, LW, "METHODOLOGY AND PROCESS FOR IMPLEMENTATION (architecture)")
tags = [("USERS", 1.62, 0.62), ("ACCESS", 2.52, 0.52), ("SERVICES", 3.34, 0.95), ("DATA", 4.52, 0.55)]
for t, y, h in tags:
    box(s, LX, y, 1.12, h, fill=NAVY, paras=[[(t, {"bold": True, "size": 10.5, "color": WHITE})]], radius=0.05)
CX, CW = 1.65, 7.05
users = [("Business / Trader", "apply, pay, renew"), ("LMO field app", "offline PWA"), ("GATC lab", "test + certify"),
         ("Controller / Admin", "dashboards, rules"), ("Citizen", "scan QR, no login")]
uw = (CW - 4 * 0.12) / 5
for i, (t, d) in enumerate(users):
    fill, ln = (LGREEN, GREEN) if i == 4 else (WHITE, BLUE)
    box(s, CX + i * (uw + 0.12), 1.62, uw, 0.62, fill=fill, line=ln, radius=0.06,
        paras=[[(t, {"bold": True, "size": 10.5, "color": NAVY})], [(d, {"size": 8.5, "color": GREY})]])
box(s, CX, 2.52, 5.55, 0.52, fill=LBLUE, line=BLUE, radius=0.05, paras=[[
    ("React PWA on Vercel  ·  Supabase Auth (phone OTP, e-Pramaan SSO later)  ·  role + district claims", {"bold": True, "size": 10, "color": NAVY})]])
box(s, CX + 5.67, 2.52, CW - 5.67, 0.52, fill=LGREEN, line=GREEN, radius=0.05, paras=[[
    ("/verify/:id", {"bold": True, "size": 10, "color": GREEN})], [("public, cached", {"size": 8.5, "color": GREY})]])
svc = [("Rule engine", "fee · validity · checklist per state", BLUE, WHITE),
       ("Workflow state machine", "submitted → scrutiny → assigned → certified", BLUE, WHITE),
       ("Assignment", "district + category → LMO / GATC", BLUE, WHITE),
       ("Edge Functions", "sign-certificate · gemini-proxy · expiry-cron", PURPLE, LPURPLE)]
sw = (CW - 3 * 0.12) / 4
for i, (t, d, c, f) in enumerate(svc):
    box(s, CX + i * (sw + 0.12), 3.34, sw, 0.95, fill=f, line=c, radius=0.06,
        paras=[[(t, {"bold": True, "size": 10.5, "color": NAVY if c == BLUE else PURPLE})], [(d, {"size": 8.5, "color": GREY})]])
data = [("Postgres + Row-Level Security", "isolation by state / district / owner"),
        ("Storage", "photos, documents, certificate PDFs"),
        ("Append-only audit log", "who · what · when, incl. AI output")]
dw = (CW - 2 * 0.12) / 3
for i, (t, d) in enumerate(data):
    box(s, CX + i * (dw + 0.12), 4.52, dw, 0.55, fill=LIGHT, line=NAVY, radius=0.05,
        paras=[[(t, {"bold": True, "size": 10, "color": NAVY})], [(d, {"size": 8.5, "color": GREY})]])
for xx in (CX + uw / 2, CX + 2 * (uw + 0.12) + uw / 2, CX + 3 * (uw + 0.12) + uw / 2):
    arrow(s, xx, 2.25, xx, 2.51, color=BLUE)
arrow(s, CX + 4 * (uw + 0.12) + uw / 2, 2.25, CX + 4 * (uw + 0.12) + uw / 2, 2.51, color=GREEN)
for xx in (CX + sw / 2, CX + 1.5 * sw + 0.12, CX + 2.5 * sw + 0.24):
    arrow(s, xx, 3.05, xx, 3.33, color=BLUE)
arrow(s, CX + 3.5 * sw + 0.36, 3.05, CX + 3.5 * sw + 0.36, 3.33, color=PURPLE)
for xx in (CX + dw / 2, CX + 1.5 * dw + 0.12, CX + 2.5 * dw + 0.24):
    arrow(s, xx, 4.3, xx, 4.51, color=NAVY)

# Gemini lane
label(s, LX, 5.18, LW, "WHERE GEMINI FITS: advisory only, the officer always decides", color=PURPLE)
g = [("Upload", "application doc or nameplate photo"), ("Gemini reads", "returns JSON: model approval no., serial, capacity"),
     ("Rule engine", "cross-checks form, fee, class"), ("Officer", "accepts or rejects each flag"),
     ("Audit log", "stores output + model version")]
gw = (LW - 4 * 0.17) / 5
for i, (t, d) in enumerate(g):
    x = LX + i * (gw + 0.17)
    c = PURPLE if i == 1 else NAVY
    box(s, x, 5.48, gw, 0.72, fill=LPURPLE if i == 1 else WHITE, line=c, lw=1.25, radius=0.06,
        paras=[[(t, {"bold": True, "size": 10.5, "color": c})], [(d, {"size": 8.5, "color": GREY})]])
    if i < 4:
        arrow(s, x + gw + 0.01, 5.84, x + gw + 0.16, 5.84, color=PURPLE)
text(s, LX, 6.3, LW, 0.55, [[("Also: ", {"bold": True, "color": PURPLE}),
     ("help chat in Hindi / Gujarati / English grounded on the Act and Rules · citizen complaint triage · plain-language pendency summaries for Controllers. ",
      {"color": INK}), ("Gemini never computes fees, MPE or pass/fail.", {"bold": True, "color": INK})]], size=9.5)

RX, RW = 8.95, 3.98
label(s, RX, 1.28, RW, "TECHNOLOGIES TO BE USED")
tech = [("React 19 · TypeScript · Vite · Tailwind", "role-based UI, installable PWA"),
        ("Workbox service worker · IndexedDB", "offline inspections, sync on reconnect"),
        ("Supabase Postgres + RLS", "one database, isolated per state / district"),
        ("Supabase Auth · Storage · Realtime · pg_cron", "OTP login, files, live dashboards, daily expiry job"),
        ("Supabase Edge Functions", "keys stay server-side, signing, Gemini proxy"),
        ("Google Gemini (Flash)", "document + photo extraction, multilingual help"),
        ("Ed25519 + SHA-256 + QR", "certificate verifiable even offline"),
        ("jsPDF · Recharts · Vercel", "certificate PDF, reports, hosting")]
for i, (t, d) in enumerate(tech):
    y = 1.6 + i * 0.47
    c = PURPLE if "Gemini" in t else NAVY
    box(s, RX, y, RW, 0.42, fill=LPURPLE if "Gemini" in t else LIGHT, align=PP_ALIGN.LEFT, margin=0.12, radius=0.05,
        paras=[[(t, {"bold": True, "size": 10.5, "color": c}), ("   " + d, {"size": 8.5, "color": GREY})]])
label(s, RX, 5.44, RW, "IMPLEMENTATION PHASES")
phases = [("1", "Prototype (live): 6 roles, full lifecycle", GREEN), ("2", "Supabase backend, real auth, RLS", BLUE),
          ("3", "Signed QR, offline PWA, Gemini copilot", PURPLE), ("4", "One-district pilot → state rollout", NAVY)]
for i, (n, t, c) in enumerate(phases):
    y = 5.76 + i * 0.27
    circle(s, RX, y, 0.23, n, fill=c, size=8.5)
    text(s, RX + 0.32, y - 0.02, RW - 0.32, 0.27, [[(t, {"size": 10, "color": INK})]])
s.notes_slide.notes_text_frame.text = (
    "Top to bottom: who uses it, how they get in, what processes the request, where it is stored. "
    "Point at the purple Edge Functions box: the Gemini key and the certificate signing key live only on the server. "
    "Then the purple lane: Gemini reads uploaded documents and nameplate photos and returns structured JSON; the rule engine "
    "cross-checks it; the officer accepts or rejects. Nothing AI-generated changes a record without a human click.")

# ================= SLIDE 4: FEASIBILITY =================
s = S[3]
team_oval(s); footer(s); remove(s, "TextBox 8")
label(s, 0.4, 1.28, 7.5, "ANALYSIS OF THE FEASIBILITY OF THE IDEA: PS requirements vs build")
rows = [("PS requirement", "Live prototype today", "Pilot build"),
        ("Stakeholder registration & profiles", "◐ 6 demo roles, profile view", "Phone OTP, e-Pramaan SSO"),
        ("Apply: verification / re-verification", "✔ 4 service types, auto fee", "Treasury / PayGov payment"),
        ("Scheduling & allocation to LMO / GATC", "✔ slots, LMO or GATC route", "District-aware auto-assign"),
        ("Inspection observations recorded", "✔ checklist + MPE table", "Readings mandatory, auto verdict"),
        ("QR digital certificate", "✔ PDF + QR + SHA-256", "Ed25519-signed, any device"),
        ("Validity tracking & expiry alerts", "✔ status + in-app alerts", "Daily job → SMS / email"),
        ("Dashboards: status, pendency, enforcement", "✔ per-role dashboards", "Realtime, state-wide"),
        ("Photo & document upload", "◐ simulated attachments", "Supabase Storage + Gemini"),
        ("Export / print certificates & reports", "✔ PDF, print, CSV", "Bulk export"),
        ("Mobile app for field officers", "◐ web, GPS, camera", "Installable offline PWA")]
tx, ty, rh = 0.4, 1.6, 0.345
cols = [(2.75, NAVY), (2.4, GREEN), (2.25, BLUE)]
for ri, row in enumerate(rows):
    y = ty + ri * rh; x = tx
    for ci, (cw_, c) in enumerate(cols):
        hdr = ri == 0
        fill = NAVY if hdr else (LIGHT if ri % 2 else WHITE)
        col = WHITE if hdr else (INK if ci == 0 else (GREEN if row[ci].startswith("✔") else (ORANGE if row[ci].startswith("◐") else NAVY)))
        box(s, x, y, cw_, rh, fill=fill, shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.LEFT, margin=0.08,
            paras=[[(row[ci], {"bold": hdr or ci == 0, "size": 9.5 if not hdr else 10, "color": col})]])
        x += cw_
text(s, 0.4, ty + 11 * rh + 0.03, 7.4, 0.25, [[("✔ working   ◐ partial in prototype, completed in pilot build", {"size": 9, "color": GREY, "italic": True})]])
feas = [("Technical", "Open-source stack. Prototype already runs the full lifecycle"),
        ("Operational", "Mirrors today's LMO / GATC process. No new roles"),
        ("Economic", "Free tiers cover a one-district pilot. No licence fees"),
        ("Legal", "LM Act 2009, General Rules 2011, GATC Rules 2013")]
fw = (7.4 - 3 * 0.13) / 4
for i, (t, d) in enumerate(feas):
    box(s, 0.4 + i * (fw + 0.13), 5.72, fw, 1.08, fill=LGREEN, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.1, radius=0.06,
        paras=[[(t, {"bold": True, "size": 11, "color": GREEN})], [(d, {"size": 9.5, "color": INK})]])

RX, RW = 8.05, 4.88
label(s, RX, 1.28, RW, "POTENTIAL CHALLENGES AND RISKS  →  STRATEGIES")
risks = [("Weak network at field sites", "Offline queue in the PWA. Signed QR verifies without internet"),
         ("Fees and validity differ by state", "Rules stored as data per state. No redeploy"),
         ("Forged or copied certificates", "Ed25519 signature, public scan, revocation list"),
         ("AI misreads a document", "Gemini only suggests. Officer confirms. Every output logged"),
         ("Personal data (DPDP Act 2023)", "RLS, minimal public fields, keys never in the browser"),
         ("Officer adoption", "Few-tap mobile flow, local languages, one-district pilot"),
         ("Payment / SMS outage", "Manual receipt entry + in-app alerts as fallback")]
for i, (r_, m) in enumerate(risks):
    y = 1.6 + i * 0.52
    box(s, RX, y, 1.95, 0.46, fill=LORANGE, shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.LEFT, margin=0.08,
        paras=[[(r_, {"bold": True, "size": 9.5, "color": ORANGE})]])
    box(s, RX + 1.95, y, RW - 1.95, 0.46, fill=LIGHT, shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.LEFT, margin=0.08,
        paras=[[(m, {"size": 9.5, "color": INK})]])
qr = pic_blob(3, "qr")
box(s, RX, 5.35, RW, 1.45, fill=LBLUE, radius=0.08)
if qr:
    s.shapes.add_picture(io.BytesIO(qr), Inches(RX + 0.12), Inches(5.44), Inches(1.27), Inches(1.27))
text(s, RX + 1.55, 5.47, RW - 1.65, 1.25, [
    [("Try the live prototype", {"bold": True, "size": 13, "color": NAVY})],
    [("Scan to open TULA. Pick any of the 6 roles and run apply → inspect → certificate → verify.", {"size": 9.5, "color": INK})],
    [("tula-legal-metrology.vercel.app", {"bold": True, "size": 10, "color": BLUE})]], spacing=3)
s.notes_slide.notes_text_frame.text = (
    "Be upfront: the prototype runs the whole lifecycle in the browser; the pilot build moves it to Supabase with real auth. "
    "The orange half-circles are honest partials. If asked about AI risk: Gemini output is a suggestion with a confidence score, "
    "never a decision, and every suggestion is audit-logged with the model version.")

# ================= SLIDE 5: IMPACT =================
s = S[4]
team_oval(s); footer(s); remove(s, "TextBox 8")
L = 0.4; W = 12.53
label(s, L, 1.28, 9, "POTENTIAL IMPACT ON THE TARGET AUDIENCE")
icons = {"LMOs & GATCs": pic_blob(4, "FaClipboard"),
         "Businesses": pic_blob(4, "FaStore"), "Consumers": pic_blob(4, "FaUsers"), "Government": pic_blob(4, "FaLandmark")}
aud = [("LMOs & GATCs", ["Daily queue sorted by district", "Readings and photos on the phone", "Certificate issued on site"], BLUE),
       ("Businesses", ["Apply and pay without a visit", "Live status of every application", "Reminder before a stamp expires"], GREEN),
       ("Consumers", ["Scan any scale or pump QR", "See if it is valid right now", "Report a problem in two taps"], ORANGE),
       ("Government", ["Pendency by district, live", "Every action in one audit trail", "Enforcement cases tracked to closure"], NAVY)]
cw = (W - 3 * 0.2) / 4
for i, (t, pts, c) in enumerate(aud):
    x = L + i * (cw + 0.2)
    box(s, x, 1.6, cw, 1.72, fill=LIGHT, radius=0.08)
    text(s, x + 0.15, 1.7, cw - 0.9, 0.35, [[(t, {"bold": True, "size": 13.5, "color": c})]])
    ic = icons.get(t)
    if ic:
        box(s, x + cw - 0.66, 1.68, 0.5, 0.5, fill=c, shape=MSO_SHAPE.OVAL)
        s.shapes.add_picture(io.BytesIO(ic), Inches(x + cw - 0.66 + 0.11), Inches(1.79), Inches(0.28), Inches(0.28))
    text(s, x + 0.15, 2.14, cw - 0.25, 1.15, [[("•  " + p_, {"size": 10.5, "color": INK})] for p_ in pts], spacing=3)

label(s, L, 3.45, 9, "PILOT KPIs: targets we will measure against today's baseline")
kpi = [("≤ 7 days", "application → certificate (median)"), ("< 3 sec", "certificate check by QR scan"),
       ("100%", "due instruments reminded 30 days early"), ("0", "paper certificates issued in pilot district")]
for i, (n, d) in enumerate(kpi):
    x = L + i * (cw + 0.2)
    box(s, x, 3.75, cw, 1.1, fill=NAVY, radius=0.08,
        paras=[[(n, {"bold": True, "size": 26, "color": WHITE})], [(d, {"size": 10, "color": "D6E2F0"})]])

label(s, L, 5.0, 11, "BENEFITS OF THE SOLUTION (SOCIAL, ECONOMIC, ENVIRONMENTAL, ETC.)")
ben = [("Social", "Fair weight at ration shops, fuel pumps and mandis. Shoppers can check for themselves.", ORANGE),
       ("Economic", "Less pendency, fewer office visits and faster renewals for traders and MSMEs.", GREEN),
       ("Environmental", "Paperless applications, certificates and records. Fewer trips to the office.", BLUE),
       ("Governance", "Transparent, auditable trail from application to enforcement for every instrument.", NAVY)]
for i, (t, d, c) in enumerate(ben):
    x = L + i * (cw + 0.2)
    box(s, x, 5.3, 1.25, 1.5, fill=c, paras=[[(t, {"bold": True, "size": 10.5, "color": WHITE})]],
        radius=0.06, margin=0.03)
    box(s, x + 1.25, 5.3, cw - 1.25, 1.5, fill=LIGHT, align=PP_ALIGN.LEFT, margin=0.12, radius=0.06,
        paras=[[(d, {"size": 10, "color": INK})]])
s.notes_slide.notes_text_frame.text = (
    "The KPI numbers are pilot targets, not measured results; say so if asked. Baseline comes from the pilot district's current "
    "register. The consumer card is the emotional hook: anyone at a petrol pump can check the dispenser in three seconds.")

# ================= SLIDE 6: RESEARCH =================
s = S[5]
team_oval(s); footer(s); remove(s, "TextBox 8")
label(s, 0.4, 1.28, 8, "DETAILS / LINKS OF THE REFERENCE AND RESEARCH WORK")
text(s, 0.95, 1.6, 3.4, 0.28, [[("RESEARCH FINDING", {"bold": True, "size": 11, "color": GREY})]])
text(s, 4.72, 1.6, 3.2, 0.28, [[("DESIGN DECISION IN TULA", {"bold": True, "size": 11, "color": GREEN})]])
res = [("Sec. 24 of the Act needs verification and stamping. GATC Rules 2013 add approved test centres as a second route.",
        "One workflow for LMO and GATC, supervised by the Controller"),
       ("Fees and validity periods differ across State Enforcement Rules.",
        "Fee, validity and checklist stored as data per state + RLS"),
       ("Instruments are re-verified every 1–2 years for their whole working life.",
        "Permanent Digital ID with full history and expiry alerts"),
       ("MPE limits come from OIML R 76 (scales) and R 117 (fuel dispensers), adopted in the Rules.",
        "Pass / fail computed from recorded readings vs class MPE"),
       ("A shopper cannot tell a genuine stamp from a fake one.",
        "Signed QR anyone can scan without login, plus complaint button")]
for i, (f, d) in enumerate(res):
    y = 1.92 + i * 0.97
    circle(s, 0.4, y + 0.25, 0.4, str(i + 1), fill=BLUE, size=11)
    box(s, 0.95, y, 3.45, 0.85, fill=LIGHT, align=PP_ALIGN.LEFT, margin=0.12, radius=0.06, paras=[[(f, {"size": 10, "color": INK})]])
    arrow(s, 4.44, y + 0.42, 4.68, y + 0.42, color=GREEN, w=2)
    box(s, 4.72, y, 3.3, 0.85, fill=LGREEN, line=GREEN, align=PP_ALIGN.LEFT, margin=0.12, radius=0.06,
        paras=[[(d, {"bold": True, "size": 10.5, "color": GREEN})]])

RX, RW = 8.3, 4.63
box(s, RX, 1.6, RW, 5.2, fill=LIGHT, radius=0.06)
text(s, RX + 0.15, 1.68, RW - 0.3, 0.35, [[("References & Links", {"bold": True, "size": 14, "color": BLUE})]])
refs = [("The Legal Metrology Act, 2009", " · consumeraffairs.gov.in/pages/legal-metrology-act"),
        ("Legal Metrology (General) Rules, 2011", " & amendments"),
        ("Legal Metrology (Government Approved Test Centre) Rules, 2013", " & amendments"),
        ("State Legal Metrology (Enforcement) Rules", ""),
        ("OIML R 76-1", " Non-automatic weighing instruments · oiml.org"),
        ("OIML R 117-1", " Dynamic measuring systems for liquids other than water"),
        ("IT Act, 2000 & DPDP Act, 2023", " electronic records, signatures, personal data"),
        ("Supabase docs", " Row Level Security, Edge Functions · supabase.com/docs"),
        ("Google Gemini API", " structured JSON output · ai.google.dev"),
        ("Working prototype", " tula-legal-metrology.vercel.app")]
tb = s.shapes.add_textbox(Inches(RX + 0.15), Inches(2.08), Inches(RW - 0.3), Inches(4.65))
tf = tb.text_frame; tf.word_wrap = True
for i, (a, b) in enumerate(refs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(7)
    r = p.add_run(); r.text = f"{i + 1}.  "; _fmt_run(r, 11, True, BLUE)
    r = p.add_run(); r.text = a; _fmt_run(r, 11, True, NAVY)
    r = p.add_run(); r.text = b; _fmt_run(r, 11, False, INK)
s.notes_slide.notes_text_frame.text = (
    "Each design choice traces to a rule or a real-world constraint. If asked about MPE: we use the accuracy-class tables "
    "the General Rules adopt from OIML R 76 / R 117; the rule engine holds them as data.")

tpl.save(OUT)
print("saved", OUT)

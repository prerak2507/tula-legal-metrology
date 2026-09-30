import io, copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SCR = sys.argv[1]
OUT = sys.argv[2]
tpl = Presentation(f"{SCR}/tpl.pptx")
old = Presentation(f"{SCR}/tula.pptx")

NAVY = "1F3F6E"; BLUE = "0070C0"; GREEN = "1E7B4A"; ORANGE = "D9661F"; PURPLE = "6B3FA0"
INK = "1F2937"; GREY = "5B6573"; LIGHT = "F1F4F8"; LBLUE = "E6EEF8"; LGREEN = "E6F4EC"
LPURPLE = "F0EAF8"; LORANGE = "FCEFE5"; WHITE = "FFFFFF"; LINE = "C9D3E0"
FONT = "Calibri"

def rgb(h): return RGBColor.from_string(h)

# ---------- helpers ----------
def _fmt_run(r, size, bold=False, color=INK, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color); r.font.name = FONT

def text(sl, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.0, italic=False, spacing=0):
    """paras: str | list of (str | list of (text, {opts}))"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    m = Inches(margin); tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    if isinstance(paras, str): paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing: para.space_after = Pt(spacing)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run(); r.text = t
            _fmt_run(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", italic))
    return tb

def box(sl, x, y, w, h, fill=LIGHT, line=None, paras=None, size=11, color=INK, bold=False,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12, lw=1.0, dash=False, margin=0.06):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.shadow.inherit = False
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            s.line.dash_style = MSO_LINE.DASH
    else: s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    m = Inches(margin); tf.margin_left = tf.margin_right = m; tf.margin_top = tf.margin_bottom = Inches(0.03)
    if paras is not None:
        if isinstance(paras, str): paras = [paras]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            runs = p if isinstance(p, list) else [(p, {})]
            for t, o in runs:
                r = para.add_run(); r.text = t
                _fmt_run(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", False))
    return s

def arrow(sl, x1, y1, x2, y2, color=BLUE, w=1.5, head=True, dash=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("h", "med")
    return c

def label(sl, x, y, w, t, color=NAVY, size=12.5):
    return text(sl, x, y, w, 0.3, [[(t, {"bold": True, "color": color, "size": size})]])

def circle(sl, x, y, d, t, fill=NAVY, size=10):
    return box(sl, x, y, d, d, fill=fill, paras=[[(t, {"bold": True, "color": WHITE, "size": size})]],
               shape=MSO_SHAPE.OVAL, margin=0)

def pic_blob(slide_idx, name_contains):
    for sh in old.slides[slide_idx].shapes:
        if sh.shape_type == 13 and any(name_contains.lower() in d.lower() for d in sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr")):
            return sh.image.blob
    for sh in old.slides[slide_idx].shapes:
        if sh.shape_type == 13 and name_contains in (sh.image.filename or "") + sh.name:
            return sh.image.blob
    for sh in old.slides[slide_idx].shapes:
        if sh.shape_type == 13 and name_contains.lower() in sh.name.lower():
            return sh.image.blob
    return None

def set_title(sl, t, size=None):
    for sh in sl.shapes:
        if sh.is_placeholder and sh.placeholder_format.type is not None and "Title" in sh.name:
            p = sh.text_frame.paragraphs[0]
            runs = p.runs
            runs[0].text = t
            if size: runs[0].font.size = Pt(size)
            for r in runs[1:]: r.text = ""
            for extra in sh.text_frame.paragraphs[1:]:
                for r in extra.runs: r.text = ""
            return sh

def team_oval(sl):
    for sh in sl.shapes:
        if sh.has_text_frame and "Team Name" in sh.text_frame.text:
            tf = sh.text_frame
            p0 = tf.paragraphs[0]
            proto = p0.runs[0]
            proto.text = "Team"; _fmt_run(proto, 10, False, NAVY)
            for r in p0.runs[1:]: r.text = ""
            p1 = tf.add_paragraph(); p1.alignment = p0.alignment
            r = p1.add_run(); r.text = "FriendlyFire"
            _fmt_run(r, 10, True, NAVY)
            for extra in tf.paragraphs[1:-1]:
                for rr in extra.runs: rr.text = ""

def footer(sl):
    for sh in sl.shapes:
        if sh.has_text_frame and "@SIH Idea submission" in sh.text_frame.text:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("Template", "TULA")

def remove(sl, name):
    for sh in list(sl.shapes):
        if sh.name == name:
            sh._element.getparent().remove(sh._element)

# ---------- drop instructions slide (7) ----------
sldIdLst = tpl.slides._sldIdLst
last = sldIdLst[6]
tpl.part.drop_rel(last.rId); sldIdLst.remove(last)

S = tpl.slides

